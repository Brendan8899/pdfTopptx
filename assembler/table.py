from pptx.util import Inches

def testForExistingTable(graphic_frame):
    if graphic_frame is None:
        return False
    else:
        return graphic_frame.has_table

def add_table_next_to_existing(slide, contentObject):
    # Define dimensions for the new table
    new_table_width = Inches(4)
    new_table_height = Inches(1.5)
    
    # Initialize variables to store the position of the existing table
    existing_table_x = None
    existing_table_y = None
    existing_table_width = None
    existing_table_height = None
    
    # Iterate through shapes to find the existing table
    for shape in slide.shapes:
        if shape.has_table:
            # Get the position and size of the existing table
            existing_table_x = shape.left
            existing_table_y = shape.top
            existing_table_width = shape.width
            existing_table_height = shape.height
            break

    if existing_table_x is not None:
        # Calculate the position for the new table
        new_table_x = existing_table_x + existing_table_width + Inches(0.5)  # Adding a 0.5-inch gap
        new_table_y = existing_table_y
        # Add the new table to the slide
        rows = len(contentObject["dataFrame"])
        cols = len(contentObject["dataFrame"][0])
        graphic_frame = slide.shapes.add_table(rows, cols, new_table_x, new_table_y, new_table_width, new_table_height)
        table = graphic_frame.table

        # Populate the new table with data
        for i in range(rows):
            for j in range(cols):
                cell = table.cell(i, j)
                cell.text = str(contentObject["dataFrame"][i][j])
    else:
        raise ValueError("No existing table found on the slide.")